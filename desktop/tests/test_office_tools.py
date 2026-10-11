import pathlib
import sys
import tempfile
import unittest

sys.path.insert(0,str(pathlib.Path(__file__).resolve().parents[1]))
from office_tools import create_spreadsheet,create_presentation

class OfficeToolsTests(unittest.TestCase):
    def test_xlsx_and_readback(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            result=create_spreadsheet(folder,"report.xlsx",[{"name":"Report","rows":[["Name","Value"],["שלום",3]]}],lambda _:True)
            self.assertIn("created_xlsx",result)
            wb=load_workbook(result["created_xlsx"])
            self.assertEqual(wb["Report"]["A2"].value,"שלום")
            self.assertTrue(wb["Report"].sheet_view.rightToLeft)
            with self.assertRaises(FileExistsError):
                create_spreadsheet(folder,"report.xlsx",[{"rows":[]}],lambda _:True)
    def test_pptx_and_readback(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as folder:
            result=create_presentation(folder,"deck.pptx",[{"title":"כותרת","bullets":["שורה א","שורה ב"]}],lambda _:True)
            slides=Presentation(result["created_pptx"]).slides
            self.assertEqual(len(slides),1)
            self.assertEqual(slides[0].shapes.title.text,"כותרת")
    def test_permission_denied(self):
        with tempfile.TemporaryDirectory() as folder:
            self.assertEqual(create_spreadsheet(folder,"report.xlsx",[{"rows":[]}],lambda _:False),{"cancelled":True})
            self.assertFalse((pathlib.Path(folder)/"report.xlsx").exists())
    def test_reject_outside_workspace(self):
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaises(ValueError):
                create_presentation(folder,"../evil.pptx",[{"title":"Oops"}],lambda _:True)

if __name__=="__main__":
    unittest.main()
