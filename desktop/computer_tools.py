"""Approved desktop UI operations; executed only after explicit user consent."""
from datetime import datetime
from pathlib import Path
from urllib.parse import quote
import subprocess
import webbrowser

def _consent(approve, explanation):
    if not approve(explanation):
        return {"cancelled": True}

def computer_action(name, args, workspace, approve):
    if not isinstance(args,dict): raise ValueError("Invalid arguments")
    if name=="screenshot":
        if (cancel:=_consent(approve,"לצלם את המסך ולשמור צילום בתיקיית העבודה? התמונה לא תישלח אוטומטית לשום שירות.")):return cancel
        import pyautogui
        folder=Path(workspace);folder.mkdir(parents=True,exist_ok=True)
        path=folder/("screenshot-"+datetime.now().strftime("%Y%m%d-%H%M%S-%f")+".png")
        pyautogui.screenshot().save(str(path))
        return {"screenshot_saved":str(path),"uploaded":False}
    if name=="click":
        x,y=args.get("x"),args.get("y")
        if type(x)!=int or type(y)!=int or not 0<=x<=10000 or not 0<=y<=10000:
            raise ValueError("Invalid coordinates")
        if (cancel:=_consent(approve,f"ללחוץ בעכבר במיקום ({x}, {y}) בחלון הפעיל?")):return cancel
        import pyautogui
        pyautogui.FAILSAFE=True
        pyautogui.click(x,y)
        return {"clicked":[x,y]}
    if name=="type_text":
        value=args.get("text")
        if not isinstance(value,str) or not 1<=len(value)<=3000:
            raise ValueError("Text length must be 1-3000")
        if (cancel:=_consent(approve,"להקליד בחלון הפעיל טקסט זה?\n\n"+value[:500])):return cancel
        import pyautogui
        import pyperclip
        pyperclip.copy(value)
        pyautogui.hotkey("ctrl","v")
        return {"typed_characters":len(value)}
    if name=="hotkey":
        keys=args.get("keys")
        permitted={("ctrl","s"),("ctrl","a"),("ctrl","c"),("ctrl","v"),("ctrl","z"),("alt","tab"),("ctrl","n"),("ctrl","o")}
        if not isinstance(keys,list) or tuple(keys) not in permitted:raise ValueError("Keyboard shortcut not allowed")
        if (cancel:=_consent(approve,"לבצע בחלון הפעיל את קיצור המקשים "+ " + ".join(keys)+"?")):return cancel
        import pyautogui
        pyautogui.hotkey(*keys)
        return {"keys":keys}
    if name=="open_browser":
        url=args.get("url","https://www.google.com")
        if not isinstance(url,str) or len(url)>1000 or not url.startswith("https://"):
            raise ValueError("Only HTTPS URLs supported")
        from urllib.parse import urlparse
        if not urlparse(url).hostname:raise ValueError("Invalid URL")
        if (cancel:=_consent(approve,"לפתוח את הכתובת הבאה בדפדפן?\n"+url)):return cancel
        webbrowser.open(url)
        return {"opened_url":url}
    if name=="create_word":
        filename=args.get("name");body=args.get("content")
        if not isinstance(filename,str) or not filename.endswith(".docx") or Path(filename).name!=filename or filename.startswith(".") or len(filename)>90:
            raise ValueError("Word filename must be a simple .docx name")
        if not isinstance(body,str) or len(body)>50000:raise ValueError("Invalid document content")
        if (cancel:=_consent(approve,"ליצור מסמך Word בשם "+filename+" בתיקיית העבודה?")):return cancel
        from docx import Document
        folder=Path(workspace);folder.mkdir(parents=True,exist_ok=True)
        path=folder/filename
        if path.is_symlink():raise ValueError("Symlink not allowed")
        if path.exists():raise FileExistsError("Document exists; choose another filename")
        doc=Document()
        for paragraph in body.splitlines():doc.add_paragraph(paragraph)
        doc.save(str(path))
        return {"created_docx":str(path)}
    if name=="project_file":
        project=args.get("project")
        relative=args.get("path")
        content=args.get("content")
        import re
        if not isinstance(project,str) or not re.fullmatch(r"[a-zA-Z0-9_ -]{1,64}",project) or project.strip()!=project:
            raise ValueError("Invalid project name")
        if not isinstance(relative,str) or len(relative)>160 or "\\\\" in relative:
            raise ValueError("Invalid project path")
        segments=relative.split("/")
        if len(segments)>5 or any(not re.fullmatch(r"[a-zA-Z0-9_. -]{1,70}",bit) or bit in (".","..") or bit.startswith(".") for bit in segments):
            raise ValueError("Invalid path components")
        if Path(relative).suffix.lower() not in {".py",".js",".ts",".tsx",".jsx",".html",".css",".json",".md",".txt",".yml",".yaml",".toml",".csv"}:
            raise ValueError("Unsupported code file extension")
        if not isinstance(content,str) or len(content)>100000:
            raise ValueError("Invalid project content")
        root=Path(workspace)/"Projects"/project
        target=root.joinpath(*segments)
        if root.is_symlink() or any(path.is_symlink() for path in [root,*target.parents] if path.exists()) or target.is_symlink():
            raise ValueError("Symlinks not allowed")
        if (cancel:=_consent(approve,"ליצור או לעדכן קובץ בפרויקט "+project+": "+relative+"?")):return cancel
        target.parent.mkdir(parents=True,exist_ok=True)
        if target.exists():
            import shutil
            backup=root/".backups"
            backup.mkdir(parents=True,exist_ok=True)
            shutil.copy2(target,backup/(datetime.now().strftime("%Y%m%d-%H%M%S-%f")+"-"+target.name))
        target.write_text(content,encoding="utf-8")
        return {"saved_project_file":str(target)}
    if name in {"workspace_files","inspect_workspace_file","open_workspace_file"}:
        import os
        root=Path(workspace).resolve()
        root.mkdir(parents=True,exist_ok=True)
        if name=="workspace_files":
            if (cancel:=_consent(approve,"להציג למודל את רשימת הקבצים והתיקיות בתיקיית העבודה?")):return cancel
            items=[]
            for file in root.rglob("*"):
                if len(items)>=150:break
                if file.is_symlink():continue
                try:
                    items.append({"path":file.relative_to(root).as_posix(),
                                  "is_dir":file.is_dir(),
                                  "size":file.stat().st_size if file.is_file() else None})
                except (OSError,ValueError):continue
            return {"workspace_items":items,"truncated":len(items)>=150}
        rel=args.get("path")
        if not isinstance(rel,str) or not rel or len(rel)>250 or "\\\\" in rel:
            raise ValueError("Invalid workspace path")
        pieces=rel.split("/")
        if any(x in ("",".","..") or x.startswith(".") for x in pieces):
            raise ValueError("Unsafe path")
        target=root.joinpath(*pieces)
        if not target.is_file() or target.is_symlink() or not target.resolve().is_relative_to(root):
            raise ValueError("File unavailable in workspace")
        ext=target.suffix.lower()
        if name=="open_workspace_file":
            safe_documents={".txt",".md",".pdf",".docx",".xlsx",".csv",".png",".jpg",".jpeg",".gif",".bmp",".webp",".rtf",".odt",".ods",".pptx"}
            if ext not in safe_documents:raise ValueError("Only non-executable document formats can be opened")
            if (cancel:=_consent(approve,"לפתוח במחשב את הקובץ "+rel+" באמצעות יישום ברירת המחדל?")):return cancel
            os.startfile(str(target))
            return {"opened_file":rel}
        if (cancel:=_consent(approve,"לקרוא מידע/תוכן מהקובץ "+rel+" ולשלוח למודל OpenAI?")):return cancel
        size=target.stat().st_size
        if size>15_000_000:raise ValueError("File exceeds inspection limit")
        result={"name":rel,"size":size,"extension":ext}
        if ext==".pdf":
            from pypdf import PdfReader
            reader=PdfReader(str(target))
            result["pages"]=len(reader.pages)
            result["text"]="\n".join((page.extract_text() or "") for page in reader.pages[:5])[:12000]
        elif ext==".docx":
            from docx import Document
            doc=Document(str(target))
            result["text"]="\n".join(paragraph.text for paragraph in doc.paragraphs)[:12000]
        elif ext in {".txt",".md",".py",".js",".json",".html",".css",".csv",".xml",".yaml",".yml",".toml",".log",".ini",".ts",".tsx",".jsx",".sql"}:
            result["text"]=target.read_text(encoding="utf-8",errors="replace")[:12000]
        else:
            result["note"]="File metadata available; content decoding not implemented for this format."
        return result
    if name=="email_draft":
        to=args.get("to");subject=args.get("subject","");body=args.get("body","")
        if not isinstance(to,str) or not 3<len(to)<255 or "@" not in to or any(c in to for c in "\r\n?&"):
            raise ValueError("Invalid email address")
        if not isinstance(subject,str) or not isinstance(body,str) or len(subject)>200 or len(body)>2000:
            raise ValueError("Invalid message")
        if (cancel:=_consent(approve,"לפתוח טיוטת דוא״ל לנמען "+to+"? לא יישלח מייל ולא יצורף צילום מסך.")):return cancel
        uri="https://mail.google.com/mail/?view=cm&fs=1&to="+quote(to)+"&su="+quote(subject)+"&body="+quote(body)
        webbrowser.open(uri)
        return {"gmail_compose_opened":to,"sent":False,"attachment_added":False,
                "note":"Attach the screenshot manually, inspect it, then send yourself."}
    raise ValueError("Unsupported desktop action")
