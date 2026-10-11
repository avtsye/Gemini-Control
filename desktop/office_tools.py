"""Office document generators confined to the GeminiControl workspace."""
import re
from pathlib import Path

def _target(workspace, filename, extension):
    if not isinstance(filename,str) or not re.fullmatch(r"[\w\- ]{1,80}\."+extension, filename, re.UNICODE):
        raise ValueError("Invalid output filename")
    root=Path(workspace).resolve()
    root.mkdir(parents=True,exist_ok=True)
    result=root/filename
    if result.is_symlink() or result.exists():
        raise FileExistsError("Document already exists; choose a new filename")
    return result

def create_spreadsheet(workspace, name, sheets, approve):
    path=_target(workspace,name,"xlsx")
    if not isinstance(sheets,list) or not 1<=len(sheets)<=10:
        raise ValueError("Expected 1-10 worksheets")
    for sheet in sheets:
        if not isinstance(sheet,dict) or not isinstance(sheet.get("rows"),list):
            raise ValueError("Invalid worksheet")
        rows=sheet["rows"]
        if len(rows)>500 or any(not isinstance(row,list) or len(row)>30 for row in rows):
            raise ValueError("Worksheet limits exceeded")
    if not approve("ליצור קובץ Excel בשם "+name+" בתיקיית העבודה?"):
        return {"cancelled":True}
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    book=Workbook()
    book.remove(book.active)
    for i,item in enumerate(sheets):
        title=str(item.get("name") or "Sheet "+str(i+1))[:31]
        if any(c in title for c in "[]:*?/\\"):
            raise ValueError("Invalid sheet name")
        sh=book.create_sheet(title)
        sh.sheet_view.rightToLeft=bool(item.get("rtl",True))
        for r,row in enumerate(item["rows"],1):
            for col,value in enumerate(row,1):
                if value is not None and not isinstance(value,(str,int,float,bool)):
                    raise ValueError("Invalid cell")
                if isinstance(value,str):
                    value=value[:5000]
                    if value.startswith(("=","+","-","@")):
                        value="'"+value
                cell=sh.cell(r,col,value)
                cell.alignment=Alignment(wrap_text=True,vertical="center")
                if r==1:
                    cell.font=Font(bold=True)
                    cell.fill=PatternFill("solid",fgColor="E8EEF7")
        sh.freeze_panes="A2"
        for col in sh.columns:
            letter=col[0].column_letter
            sh.column_dimensions[letter].width=min(48,max(12,max((len(str(c.value or "")) for c in col),default=12)+2))
    book.save(path)
    return {"created_xlsx":str(path),"sheets":len(sheets)}

def create_presentation(workspace,name,slides,approve):
    path=_target(workspace,name,"pptx")
    if not isinstance(slides,list) or not 1<=len(slides)<=40:
        raise ValueError("Expected 1-40 slides")
    for slide in slides:
        if not isinstance(slide,dict) or not isinstance(slide.get("title",""),str):
            raise ValueError("Invalid slide")
    if not approve("ליצור מצגת PowerPoint בשם "+name+" בתיקיית העבודה?"):
        return {"cancelled":True}
    from pptx import Presentation
    from pptx.util import Inches,Pt
    presentation=Presentation()
    presentation.slide_width=Inches(13.333)
    presentation.slide_height=Inches(7.5)
    for item in slides:
        slide=presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text=item.get("title","")[:250]
        contents=item.get("bullets",[])
        if not isinstance(contents,list) or len(contents)>12 or any(not isinstance(v,str) for v in contents):
            raise ValueError("Invalid bullets")
        frame=slide.placeholders[1].text_frame
        frame.clear()
        for j,value in enumerate(contents):
            p=frame.paragraphs[0] if j==0 else frame.add_paragraph()
            p.text=value[:700]
            p.font.size=Pt(24)
        if not contents:
            frame.text=""
    presentation.save(path)
    return {"created_pptx":str(path),"slides":len(slides)}
